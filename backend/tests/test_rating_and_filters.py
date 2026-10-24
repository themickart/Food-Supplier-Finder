from io import BytesIO

import pytest
from openpyxl import load_workbook


def test_information_index_is_explained_and_not_a_reliability_rating(client):
    item = client.get("/api/v1/suppliers/demo-01?mode=demo").json()
    rating = item["rating"]
    assert rating["value"] == 95
    assert sum(criterion["max"] for criterion in rating["criteria"]) == 100
    assert sum(criterion["points"] for criterion in rating["criteria"]) - rating["conflict_penalty"] == rating["value"]
    assert "не надёжность" not in rating["label"]  # The label describes information, never company performance.
    assert "а не надёжность компании" in rating["disclaimer"]
    conflict = client.get("/api/v1/suppliers/demo-06?mode=demo").json()["rating"]
    assert conflict["conflict_penalty"] == 20
    stale = client.get("/api/v1/suppliers/demo-08?mode=demo").json()["rating"]
    assert next(c for c in stale["criteria"] if c["id"] == "freshness")["points"] == 0


def test_source_alone_does_not_get_high_score_or_verification(client):
    result = client.post(
        "/api/v1/suppliers/import",
        json={"name": "Минимальная тестовая запись", "source_url": "https://test-source.ru/", "category": "dairy"},
    ).json()["supplier"]
    assert result["rating"]["value"] == 20
    assert result["status"] == "unverified"
    assert result["checked_at"] is None


def test_score_sorting_threshold_and_status_filter(client):
    response = client.get("/api/v1/suppliers", params={"mode": "demo", "sort": "rating", "min_rating": 75})
    assert response.status_code == 200, response.text
    items = response.json()["items"]
    values = [item["rating"]["value"] for item in items]
    assert values == sorted(values, reverse=True)
    assert all(value >= 75 for value in values)
    review = client.get("/api/v1/suppliers", params={"mode": "demo", "status": "needs_review"}).json()["items"]
    assert review and all(item["status"] == "needs_review" for item in review)
    assert client.get("/api/v1/suppliers", params={"mode": "real", "min_rating": 75}).json()["total"] == 0


@pytest.mark.parametrize("unit,lower,upper,expected", [("кг", 500, 700, "demo-02"), ("уп.", 1000, 1500, "demo-04")])
def test_price_range_only_compares_same_product_currency_and_unit(client, unit, lower, upper, expected):
    params = {"mode": "demo", "price_product": "моцарелла", "price_unit": unit, "price_from": lower, "price_to": upper}
    response = client.get("/api/v1/suppliers", params=params)
    assert response.status_code == 200, response.text
    assert [item["id"] for item in response.json()["items"]] == [expected]
    assert client.get("/api/v1/suppliers", params={**params, "price_currency": "USD"}).json()["total"] == 0


def test_unknown_price_is_not_zero_and_order_units_are_not_mixed(client):
    items = client.get(
        "/api/v1/suppliers", params={"mode": "demo", "price_product": "клубника", "price_unit": "кг", "price_from": 0}
    ).json()["items"]
    assert {item["id"] for item in items} == {"demo-05", "demo-06"}
    items = client.get("/api/v1/suppliers", params={"mode": "demo", "order_unit": "₽", "order_to": 5000}).json()[
        "items"
    ]
    assert {item["id"] for item in items} == {"demo-03", "demo-10"}
    assert (
        client.get("/api/v1/suppliers", params={"mode": "demo", "order_unit": "$", "order_to": 5000}).json()["total"]
        == 0
    )


@pytest.mark.parametrize(
    "params",
    [
        {"price_from": 1},
        {"price_from": 100, "price_to": 10, "price_product": "сыр", "price_unit": "кг"},
        {"order_from": 1},
        {"min_rating": 101},
    ],
)
def test_unsafe_or_ambiguous_filter_ranges_are_validation_errors(client, params):
    response = client.get("/api/v1/suppliers", params={"mode": "demo", **params})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == "validation"


def test_price_and_order_conditions_must_belong_to_same_term(client, real_supplier):
    data = {
        **real_supplier,
        "terms": [
            {"product": "Моцарелла", "price": 500, "unit": "кг", "currency": "RUB"},
            {"product": "Моцарелла", "min_order": 10, "min_order_unit": "кг", "unit": "кг"},
        ],
    }
    assert client.post("/api/v1/suppliers/import", json=data).status_code == 201
    response = client.get(
        "/api/v1/suppliers",
        params={"price_product": "моцарелла", "price_unit": "кг", "price_to": 600, "order_unit": "кг", "order_to": 20},
    )
    assert response.status_code == 200
    assert response.json()["total"] == 0


def test_extended_import_edit_and_source_history(client, real_supplier):
    original = client.post("/api/v1/suppliers/import", json=real_supplier).json()["supplier"]
    data = {
        **real_supplier,
        "city": "Казань",
        "delivery_regions": ["Казань", "Москва"],
        "terms": [
            {
                "product": "Моцарелла",
                "price": 600,
                "unit": "кг",
                "currency": "RUB",
                "packaging": "Пакет 2 кг",
                "volume": "От 20 кг",
                "min_order": 20,
                "min_order_unit": "кг",
            }
        ],
        "documents": [{"name": "Тестовый документ", "url": "https://supplier-test.ru/document"}],
    }
    response = client.patch(f"/api/v1/suppliers/{original['id']}", json=data)
    assert response.status_code == 200, response.text
    updated = response.json()["supplier"]
    assert updated["id"] == original["id"]
    assert updated["city"] == "Казань"
    assert updated["terms"][0]["price"] == 600
    assert updated["terms"][0]["checked_at"] is None
    assert updated["sources"][-1]["previous"]["city"] == "Екатеринбург"
    assert updated["sources"][0]["claims"]["city"] == "Екатеринбург"
    assert updated["status"] == "unverified"
    assert client.get("/api/v1/suppliers").json()["total"] == 1
    assert client.patch("/api/v1/suppliers/demo-01", json=data).status_code == 404
    filtered = client.get(
        "/api/v1/suppliers",
        params={"price_product": "моцарелла", "price_unit": "кг", "price_from": 590, "commercial": True},
    ).json()
    assert filtered["total"] == 1
    exported = client.post(
        "/api/v1/export",
        json={"format": "xlsx", "filters": {"price_product": "моцарелла", "price_unit": "кг", "price_from": 590}},
    )
    sheet = load_workbook(BytesIO(exported.content)).active
    assert sheet.max_row == 2
    headers = [cell.value for cell in sheet[1]]
    assert "Индекс данных / 100" in headers
    assert "а не надёжность компании" in sheet.cell(2, headers.index("Что означает индекс") + 1).value


def test_edit_does_not_silently_move_record_onto_another_supplier_domain(client, real_supplier):
    a = client.post("/api/v1/suppliers/import", json=real_supplier).json()["supplier"]
    b = client.post(
        "/api/v1/suppliers/import",
        json={**real_supplier, "name": "Другая тестовая запись", "website": "https://another-test.ru"},
    ).json()["supplier"]
    assert client.patch(f"/api/v1/suppliers/{b['id']}", json=real_supplier).status_code == 409
    assert client.get(f"/api/v1/suppliers/{b['id']}").json()["website"] != a["website"]


def test_overview_and_match_are_scoped_and_explained(client):
    assert client.get("/api/v1/overview").json()["total"] == 0
    demo = client.get("/api/v1/overview?mode=demo").json()
    assert demo["total"] == 12 and len(demo["sources"]) == 12
    assert sum(demo["categories"].values()) == 12
    items = client.get("/api/v1/suppliers", params={"mode": "demo", "q": "моцарелла", "region": "Екатеринбург"}).json()[
        "items"
    ]
    assert items and all(item["match"]["value"] == 100 for item in items)
    assert all("Запрос найден" in item["match"]["reasons"][0] for item in items)


def test_matching_commercial_term_is_preserved_in_results_and_comparison(client, real_supplier):
    data = {
        **real_supplier,
        "products": [],
        "terms": [
            {"product": "Молоко", "price": 90, "unit": "л", "currency": "RUB"},
            {"product": "Моцарелла", "price": 600, "unit": "кг", "currency": "RUB"},
        ],
    }
    item = client.post("/api/v1/suppliers/import", json=data).json()["supplier"]
    assert "Моцарелла" in item["products"]
    text_search = client.get("/api/v1/suppliers", params={"q": "моцарелла"}).json()
    assert text_search["total"] == 1
    assert text_search["items"][0]["matched_term_index"] == 1
    filters = {"price_product": "моцарелла", "price_unit": "кг", "price_from": 500}
    result = client.get("/api/v1/suppliers", params=filters).json()["items"][0]
    assert result["matched_term_index"] == 1
    comparison = client.post("/api/v1/compare", json={"ids": [item["id"]], "filters": filters}).json()["items"][0]
    assert comparison["matched_term_index"] == 1
    assert comparison["terms"][comparison["matched_term_index"]]["price"] == 600
    assert (
        client.post(
            "/api/v1/compare", json={"ids": [item["id"]], "mode": "real", "filters": {"mode": "demo"}}
        ).status_code
        == 422
    )
