import csv
from io import BytesIO, StringIO

import pytest
from openpyxl import load_workbook


def test_verified_contact_filter_does_not_treat_manual_or_demo_as_checked(client, real_supplier):
    client.post("/api/v1/suppliers/import", json=real_supplier)
    assert client.get("/api/v1/suppliers", params={"contacts": True}).json()["total"] == 1
    assert client.get("/api/v1/suppliers", params={"verified_contacts": True}).json()["total"] == 0
    assert client.get("/api/v1/suppliers", params={"mode": "demo", "verified_contacts": True}).json()["total"] == 0


def test_plain_search_includes_delivery_geography(client):
    response = client.get("/api/v1/suppliers", params={"mode": "demo", "q": "Казань Екатеринбург"})
    assert response.status_code == 200
    assert [s["name"] for s in response.json()["items"]] == ["Фермерская линия"]


def test_demo_is_never_in_real_catalog_or_details(client):
    assert client.get("/api/v1/suppliers").json()["total"] == 0
    assert client.get("/api/v1/suppliers", params={"mode": "demo"}).json()["total"] == 12
    assert client.get("/api/v1/suppliers/demo-01").status_code == 404
    assert client.post("/api/v1/compare", json={"ids": ["demo-01"]}).status_code == 404
    assert client.get("/api/v1/meta").json()["counts"] == {"real": 0, "demo": 12}


def test_filters_russian_case_units_and_delivery_geography(client):
    response = client.get("/api/v1/suppliers", params={"mode": "demo", "q": "МОЦАРЕЛЛА", "region": "Екатеринбург"})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 2
    response = client.get("/api/v1/suppliers", params={"mode": "demo", "category": "packaging", "documents": "true"})
    assert response.json()["total"] == 2
    delivery = client.get(
        "/api/v1/suppliers", params={"mode": "demo", "category": "dairy", "delivery_region": "Екатеринбург"}
    ).json()
    assert delivery["total"] == 4
    assert {item["city"] for item in delivery["items"]} == {"Екатеринбург", "Казань", "Москва"}


def test_empty_query_errors_partial_external_and_no_ai(client):
    assert client.post("/api/v1/search", json={"filters": {}}).status_code == 422
    search = client.post(
        "/api/v1/search", json={"filters": {"q": "моцарелла", "mode": "demo"}, "external": True}
    ).json()
    assert search["status"] == "partial"
    assert search["external"]["status"] == "not_configured"
    assert search["total"] == 3
    assert search["ai"] == "disabled"
    health = client.get("/api/v1/health").json()
    assert health["paid_requests"] is False
    assert client.get("/api/v1/suppliers", params={"q": "несуществующийтовар", "mode": "demo"}).json()["items"] == []


def test_manual_import_deduplication_and_provenance(client, real_supplier):
    first = client.post("/api/v1/suppliers/import", json=real_supplier)
    assert first.status_code == 201, first.text
    item = first.json()["supplier"]
    assert first.json()["created"] is True
    assert item["status"] == "unverified"
    assert item["checked_at"] is None
    assert item["sources"][0]["retrieved_at"]
    assert item["sources"][0]["checked_at"] is None
    assert item["contacts"][0]["verification"] == "user_supplied"
    again = client.post(
        "/api/v1/suppliers/import", json={**real_supplier, "website": "https://supplier-test.ru/"}
    ).json()
    assert again["created"] is False
    assert again["matched_by"] == "domain"
    assert again["supplier"]["id"] == item["id"]
    assert client.get("/api/v1/suppliers").json()["total"] == 1
    third = client.post("/api/v1/suppliers/import", json=real_supplier).json()
    assert len(third["supplier"]["sources"]) == 2


def test_conflicting_values_preserved_in_sources(client, real_supplier):
    original = client.post("/api/v1/suppliers/import", json=real_supplier).json()["supplier"]
    updated = client.post(
        "/api/v1/suppliers/import",
        json={**real_supplier, "city": "Москва", "source_url": "https://supplier-test.ru/new"},
    ).json()["supplier"]
    assert updated["id"] == original["id"]
    assert updated["city"] == "Екатеринбург"
    assert updated["status"] == "needs_review"
    assert updated["sources"][1]["claims"]["city"] == "Москва"
    assert updated["sources"][0]["claims"]["city"] == "Екатеринбург"


def test_ambiguous_domain_does_not_merge_different_legal_entities(client, real_supplier):
    a = client.post("/api/v1/suppliers/import", json={**real_supplier, "inn": "7707083893"}).json()
    b = client.post(
        "/api/v1/suppliers/import", json={**real_supplier, "name": "Другая тестовая компания", "inn": "7736207543"}
    ).json()
    assert a["created"] and b["created"]
    assert a["supplier"]["id"] != b["supplier"]["id"]
    assert b["supplier"]["conflicts"]
    assert client.get("/api/v1/suppliers").json()["total"] == 2


@pytest.mark.parametrize(
    "url",
    [
        "file:///etc/passwd",
        "http://localhost/test",
        "http://127.0.0.1/test",
        "http://169.254.169.254/meta",
        "https://user:password@valid.ru",
        "javascript:alert(1)",
        "https://valid.ru:9000",
        "https://fake.example",
    ],
)
def test_unsafe_source_urls_rejected_without_network(client, real_supplier, url):
    result = client.post("/api/v1/suppliers/import", json={**real_supplier, "source_url": url})
    assert result.status_code == 422
    assert result.json()["error"]["code"] == "validation"


def test_validation_and_cross_origin_writes(client, real_supplier):
    assert client.post("/api/v1/suppliers/import", json={**real_supplier, "inn": "1234567890"}).status_code == 422
    assert client.post("/api/v1/suppliers/import", json={**real_supplier, "is_demo": True}).status_code == 422
    assert (
        client.post("/api/v1/suppliers/import", json=real_supplier, headers={"origin": "https://evil.test"}).status_code
        == 403
    )


def test_comparison_limit_missing_data_and_explicit_units(client):
    assert (
        client.post("/api/v1/compare", json={"ids": [f"demo-{i:02}" for i in range(1, 7)], "mode": "demo"}).status_code
        == 422
    )
    result = client.post("/api/v1/compare", json={"ids": ["demo-02", "demo-04", "demo-12"], "mode": "demo"}).json()
    assert result["items"][0]["terms"][0]["unit"] == "кг"
    assert result["items"][1]["terms"][0]["unit"] == "уп."
    assert result["items"][2]["terms"] == []
    assert "победитель" not in result
    assert client.get("/api/v1/suppliers/demo-08?mode=demo").json()["status"] == "needs_review"


def test_pagination_deterministic_sort_and_favorite_collection(client):
    first = client.get("/api/v1/suppliers", params={"mode": "demo", "page_size": 5, "sort": "name"}).json()
    second = client.get("/api/v1/suppliers", params={"mode": "demo", "page_size": 5, "sort": "name", "page": 2}).json()
    assert len(first["items"]) == 5
    assert not {s["id"] for s in first["items"]} & {s["id"] for s in second["items"]}
    filtered = client.post(
        "/api/v1/collection", json={"ids": ["demo-01", "demo-05"], "filters": {"mode": "demo", "category": "frozen"}}
    ).json()
    assert filtered["total"] == 1
    assert filtered["items"][0]["id"] == "demo-05"


def test_csv_selected_rows_quotes_and_formula_injection(client, real_supplier):
    supplier = client.post(
        "/api/v1/suppliers/import",
        json={**real_supplier, "name": '=HYPERLINK("malicious")', "description": "Строка, с запятой\nи переносом"},
    ).json()["supplier"]
    response = client.post(
        "/api/v1/export",
        json={"format": "csv", "ids": [supplier["id"]], "include_notes": True, "notes": {supplier["id"]: "  =1+1"}},
    )
    assert response.status_code == 200
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
    assert len(rows) == 2
    assert rows[1][0].startswith("'=")
    assert rows[1][-1] == "'  =1+1"
    assert "https://supplier-test.ru/contacts" in rows[1]
    # Notes are request-only and never persist in the supplier record.
    detail = client.get(f"/api/v1/suppliers/{supplier['id']}").json()
    assert "notes" not in detail


def test_xlsx_filtered_rows_demo_labels_and_no_formulas(client):
    response = client.post(
        "/api/v1/export",
        json={
            "format": "xlsx",
            "filters": {"mode": "demo", "category": "packaging", "region": "Казань"},
            "include_notes": True,
            "notes": {"demo-10": "@SUM(A1)"},
        },
    )
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet.max_row == 2
    assert sheet.cell(2, 1).value == "Бережная упаковка"
    assert "ДЕМО" in sheet.cell(2, 2).value
    assert sheet.cell(2, sheet.max_column).value == "'@SUM(A1)"
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert sheet.freeze_panes == "C2"


def test_export_rejects_cross_mode_and_invalid_format(client):
    assert client.post("/api/v1/export", json={"format": "xlsx", "ids": ["demo-01"]}).status_code == 404
    assert client.post("/api/v1/export", json={"format": "pdf"}).status_code == 422
    empty = client.post("/api/v1/export", json={"format": "csv"})
    assert len(list(csv.reader(StringIO(empty.content.decode("utf-8-sig")), delimiter=";"))) == 1
