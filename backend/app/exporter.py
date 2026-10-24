import csv
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .catalog import CATEGORIES, serialize

HEADERS = [
    "Название",
    "Тип данных",
    "Категории",
    "Товары",
    "Город",
    "Регион",
    "Адрес",
    "Описание",
    "Сайт",
    "ИНН",
    "Телефон",
    "Email",
    "Условия закупки",
    "Доставка",
    "Регионы доставки",
    "Документы",
    "Источники",
    "Дата получения",
    "Последняя проверка",
    "Статус",
    "Противоречия",
    "Индекс данных / 100",
    "Что означает индекс",
]


def safe_cell(value) -> str:
    value = "" if value is None else str(value)
    value = "".join(c for c in value if c in "\t\r\n" or ord(c) >= 32)
    # Excel and LibreOffice can interpret formulas after leading whitespace/control characters.
    if value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        value = "'" + value
    return value[:32000]


def rows_for(suppliers, notes, include_notes):
    rows = [HEADERS + (["Личная заметка"] if include_notes else [])]
    for supplier in suppliers:
        item = serialize(supplier)
        terms = []
        for term in supplier.terms:
            price = (
                f"{term['price']} {term.get('currency', 'RUB')}/{term.get('unit', 'Не указано')}"
                if term.get("price") is not None
                else "Цена не указана"
            )
            minimum = (
                f"{term['min_order']} {term.get('min_order_unit', '')}"
                if term.get("min_order") is not None
                else "Не указано"
            )
            terms.append(
                f"{term.get('product', 'Товар не указан')}: {price}; фасовка: {term.get('packaging') or 'Не указано'}; объём: {term.get('volume') or 'Не указано'}; мин. заказ: {minimum}; источник: {term.get('source_url') or 'Не указано'}; проверка: {term.get('checked_at') or 'Не проводилась'}"
            )
        row = [
            supplier.name,
            "ДЕМО — вымышленная компания" if supplier.is_demo else "Реальная запись, требует проверки",
            ", ".join(c["name"] for c in CATEGORIES if c["id"] in supplier.categories),
            ", ".join(supplier.products),
            supplier.city,
            supplier.region,
            supplier.address,
            supplier.description,
            supplier.website,
            supplier.inn,
            ", ".join(c["value"] for c in supplier.contacts if c["kind"] == "phone"),
            ", ".join(c["value"] for c in supplier.contacts if c["kind"] == "email"),
            "\n".join(terms),
            supplier.delivery_info,
            ", ".join(supplier.delivery_regions),
            "\n".join(f"{d['name']}: {d.get('url', '')}" for d in supplier.documents),
            "\n".join(s["url"] for s in supplier.sources),
            "\n".join(s["retrieved_at"] for s in supplier.sources),
            item["checked_at"] or "Не проводилась",
            {
                "needs_review": "Требует перепроверки",
                "unverified": "Не проверено",
                "source_checked": "Проверено по источнику",
            }[item["status"]],
            "\n".join(supplier.conflicts),
            item["rating"]["value"],
            item["rating"]["disclaimer"],
        ]
        if include_notes:
            row.append(notes.get(supplier.id, ""))
        rows.append([safe_cell(value) for value in row])
    return rows


def export_file(suppliers, format, notes, include_notes):
    rows = rows_for(suppliers, notes, include_notes)
    if format == "csv":
        stream = StringIO(newline="")
        writer = csv.writer(stream, delimiter=";", quoting=csv.QUOTE_ALL)
        writer.writerows(rows)
        return stream.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Поставщики"
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="284F46")
        cell.font = Font(color="FFFFFF", bold=True)
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 30
        for cell in column:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
